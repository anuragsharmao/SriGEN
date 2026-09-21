"""SriGEN backend test suite."""

import asyncio
import json

import pytest

from app.core.config import settings
from app.db.models import DeliverableDraftModel, DisclosureItemModel
from app.db.schemas import (
    AudienceType,
    CommunicationObjective,
    ContentStyle,
    DeliverableType,
    DetailFocus,
    GenerateRequest,
    LanguageType,
    LengthType,
    ToneType,
)
from app.services.fact_graph import fact_graph_service
from app.services.grounding_guard import grounding_guard
from app.services.ingestion import UploadTooLargeError, ingest_document
from app.services.resolver import resolver
from app.services.sensitivity_firewall import firewall

pytestmark = pytest.mark.filterwarnings("ignore::DeprecationWarning")

SAMPLE_SOURCE = """
An operational incident was detected at Sector-7 on 15 August, affecting 24 systems.
Unit 4821 responded within the hour and restored all affected systems by evening.
No further disruption has been reported since containment was completed.
"""


# ---------------------------------------------------------------------------
# Phase 0: credential/config hygiene, LLM failure handling, entity detection
# ---------------------------------------------------------------------------

def test_no_hardcoded_demo_values_in_app_code():
    """Phase 0.5 acceptance check: zero hardcoded proper nouns in application code."""
    import subprocess
    result = subprocess.run(
        ["grep", "-rliE", "--include=*.py", "kanpur|lucknow|vikram|varanasi|zulu", "app/"],
        capture_output=True, text=True,
    )
    assert result.returncode != 0, f"Found hardcoded demo values in app/: {result.stdout}"


def test_requirements_txt_is_parseable():
    """Phase 0 regression check: requirements.txt must not contain literal backslash-n corruption."""
    content = open("requirements.txt", encoding="utf-8").read()
    assert "\\n" not in content
    lines = [l for l in content.splitlines() if l.strip() and not l.startswith("#")]
    assert any("en_core_web_sm" in l for l in lines)


def test_llm_unavailable_returns_503(client, fake_llm):
    """Phase 0.2 flagship acceptance test: with the LLM unavailable, /api/generate
    must return 503 with LLM_UNAVAILABLE, and must NEVER return 200 with content."""
    fake_llm.should_fail = True
    res = client.post("/api/generate", json={
        "source_text": SAMPLE_SOURCE,
        "deliverable_types": ["executive_summary"],
    })
    assert res.status_code == 503
    body = res.json()
    assert body["detail"]["error"] == "LLM_UNAVAILABLE"


def test_health_reports_llm_reachable(client, fake_llm):
    fake_llm.should_fail = False
    res = client.get("/health")
    assert res.status_code == 200
    assert res.json()["llm_reachable"] is True

    fake_llm.should_fail = True
    res2 = client.get("/health")
    assert res2.json()["llm_reachable"] is False
    assert res2.json()["status"] == "degraded"


def test_entity_boilerplate_and_sentence_initial_not_flagged():
    """Phase 0.3: ordinary capitalized words (sentence-initial, boilerplate) must
    not be treated as named entities by the regex fallback path."""
    text = "Operational readiness remains high. Immediate action was taken. Reaffirms commitment to safety."
    entities = grounding_guard.extract_named_entities(text)
    for boilerplate_word in ["operational", "immediate", "reaffirms"]:
        assert boilerplate_word not in entities


# ---------------------------------------------------------------------------
# Phase 0.3/0.4: grounding acceptance test + placeholder reading-order numbering
# ---------------------------------------------------------------------------

def test_faithful_summary_scores_high_and_tampered_summary_scores_low():
    """Phase 0.3 acceptance test: a faithful summary scores >= 90; a version with
    one number changed scores < 70 and the mismatch names the specific number."""
    source = (
        "An incident affected 24 systems at the facility. Response teams restored "
        "all systems by evening. No further disruption has been reported since containment."
    )
    faithful_summary = (
        "The incident affected 24 systems at the facility. "
        "Response teams restored all systems by evening. "
        "No further disruption has been reported since containment."
    )
    tampered_summary = (
        "The incident affected 42 systems at the facility. "
        "Response teams restored all systems by evening. "
        "No further disruption has been reported since containment."
    )

    faithful_score, faithful_claims = asyncio.run(
        grounding_guard.verify_content(content=faithful_summary, source_text=source)
    )
    tampered_score, tampered_claims = asyncio.run(
        grounding_guard.verify_content(content=tampered_summary, source_text=source)
    )

    assert faithful_score >= 90.0
    assert tampered_score < 70.0

    all_mismatch_descriptions = " ".join(
        m.description for c in tampered_claims for m in c.entity_mismatches
    )
    assert "42" in all_mismatch_descriptions


def test_placeholder_numbering_follows_reading_order_not_replacement_order():
    """Phase 0.4 acceptance test."""
    text = "Incident occurred near Sector-7. Backup response staged at Sector-9. Sector-7 remains the primary site."
    redacted, actions, placeholder_map, degraded = asyncio.run(firewall.apply_redaction(text, audience=AudienceType.GENERAL_PUBLIC))

    assert "[LOCATION_1]" in redacted
    assert "[LOCATION_2]" in redacted
    first_idx = redacted.index("[LOCATION_1]")
    second_idx = redacted.index("[LOCATION_2]")
    assert first_idx < second_idx
    assert placeholder_map["[LOCATION_1]"].lower() == "sector-7"
    assert placeholder_map["[LOCATION_2]"].lower() == "sector-9"
    # Second mention of Sector-7 reuses [LOCATION_1], not a new number.
    assert redacted.count("[LOCATION_1]") == 2


# ---------------------------------------------------------------------------
# Phase 1: content-only structured deliverables (Presentation / Infographic / Video)
# ---------------------------------------------------------------------------

def test_presentation_generates_structured_content(client):
    res = client.post("/api/generate", json={
        "source_text": SAMPLE_SOURCE,
        "deliverable_types": ["presentation"],
    })
    assert res.status_code == 200
    draft = res.json()["drafts"][0]
    assert draft["deliverable_type"] == "presentation"
    structured = draft["structured_content"]
    assert structured is not None
    assert len(structured["slides"]) >= 1
    assert structured["slides"][0]["layout"] == "title"
    assert structured["slides"][-1]["layout"] == "closing"
    for slide in structured["slides"]:
        assert slide["speaker_notes"].strip() != ""
    assert "Slide" in draft["draft_content"]


def test_infographic_generates_structured_content(client):
    res = client.post("/api/generate", json={
        "source_text": SAMPLE_SOURCE,
        "deliverable_types": ["infographic"],
    })
    assert res.status_code == 200
    draft = res.json()["drafts"][0]
    structured = draft["structured_content"]
    assert structured is not None
    assert structured["layout_recommendation"].strip() != ""
    assert len(structured["sections"]) >= 1


def test_video_package_covers_all_six_ps_artefacts_with_exact_timing(client):
    res = client.post("/api/generate", json={
        "source_text": SAMPLE_SOURCE,
        "deliverable_types": ["video_package"],
        "length": "Standard",
    })
    assert res.status_code == 200
    draft = res.json()["drafts"][0]
    structured = draft["structured_content"]
    assert structured is not None

    assert structured["narration_script"].strip() != ""
    assert len(structured["scenes"]) >= 1
    assert len(structured["subtitles"]) >= 1
    assert len(structured["visual_recommendations"]) >= 1

    assert structured["total_duration_seconds"] == 60.0

    scenes = sorted(structured["scenes"], key=lambda s: s["scene_number"])
    assert scenes[0]["start_seconds"] == 0.0
    assert scenes[-1]["end_seconds"] == 60.0
    for a, b in zip(scenes, scenes[1:]):
        assert a["end_seconds"] == b["start_seconds"]

    cues = structured["subtitles"]
    for i, cue in enumerate(cues, start=1):
        assert cue["index"] == i
        lines = cue["text"].split("\n")
        assert len(lines) <= 2
        assert all(len(line) <= 42 for line in lines)
        duration = cue["end_seconds"] - cue["start_seconds"]
        assert 1.2 - 1e-6 <= duration <= 6.0 + 1e-6
    for a, b in zip(cues, cues[1:]):
        assert a["end_seconds"] <= b["start_seconds"]


def test_video_package_brief_and_detailed_durations():
    from app.adapters.video_package import VideoPackageAdapter

    adapter = VideoPackageAdapter()
    fact_graph = asyncio.run(fact_graph_service.extract_fact_graph(SAMPLE_SOURCE))

    for length, expected in [(LengthType.BRIEF, 30.0), (LengthType.DETAILED, 120.0)]:
        spec = resolver.resolve_spec(DeliverableType.VIDEO_PACKAGE, GenerateRequest(
            source_text=SAMPLE_SOURCE, deliverable_types=[DeliverableType.VIDEO_PACKAGE], length=length,
        ))
        package = asyncio.run(adapter.generate_structured(fact_graph, spec, SAMPLE_SOURCE))
        assert package.total_duration_seconds == expected


def test_export_json_for_structured_deliverable_and_415_for_plain_text(client):
    res = client.post("/api/generate", json={
        "source_text": SAMPLE_SOURCE,
        "deliverable_types": ["infographic", "executive_summary"],
    })
    assert res.status_code == 200
    drafts = res.json()["drafts"]
    infographic_draft = next(d for d in drafts if d["deliverable_type"] == "infographic")
    text_draft = next(d for d in drafts if d["deliverable_type"] == "executive_summary")

    ok = client.get(f"/api/dashboard/draft/{infographic_draft['id']}/export/json")
    assert ok.status_code == 200
    assert "sections" in ok.json()["structured_content"]

    bad = client.get(f"/api/dashboard/draft/{text_draft['id']}/export/json")
    assert bad.status_code == 415


def test_custom_deliverable_uses_additional_instructions_as_format_brief(client):
    res = client.post("/api/generate", json={
        "source_text": SAMPLE_SOURCE,
        "deliverable_types": ["custom"],
        "additional_instructions": "Produce a haiku-length internal Slack update.",
    })
    assert res.status_code == 200
    draft = res.json()["drafts"][0]
    assert draft["deliverable_type"] == "custom"
    assert draft["draft_content"].strip() != ""


def test_multiselect_shares_one_source_and_one_fact_graph(client):
    res = client.post("/api/generate", json={
        "source_text": SAMPLE_SOURCE,
        "deliverable_types": ["executive_summary", "linkedin_post", "advisory"],
    })
    assert res.status_code == 200
    body = res.json()
    assert len(body["drafts"]) == 3
    source_id = body["source_id"]
    from app.db.database import SessionLocal
    db = SessionLocal()
    try:
        drafts = db.query(DeliverableDraftModel).filter(DeliverableDraftModel.batch_id == body["batch_id"]).all()
        assert len(drafts) == 3
        assert all(d.source_id == source_id for d in drafts)
    finally:
        db.close()


# ---------------------------------------------------------------------------
# Phase 2: disclosure control — grouped/individual/bulk decisions, dual-scan, gate
# ---------------------------------------------------------------------------

def _generate_with_sensitive_content(client):
    text = (
        "An incident occurred at Sector-7. Unit 4821 responded. "
        "Operation Falcon was activated to contain the breach. "
        "Col. Sharma authorized the response from the regional office."
    )
    res = client.post("/api/generate", json={
        "source_text": text,
        "deliverable_types": ["executive_summary"],
    })
    assert res.status_code == 200
    return res.json()


def test_disclosure_items_created_from_output_side_scan(client):
    body = _generate_with_sensitive_content(client)
    draft_id = body["drafts"][0]["id"]
    review = client.get(f"/api/disclosure/draft/{draft_id}/review")
    assert review.status_code == 200
    data = review.json()
    assert data["total_items"] >= 1
    for group in data["groups"]:
        for item in group["items"]:
            assert item["detected_at"] in ("output", "both")
            assert item["analyst_choice"] is None
            assert item["suggested_default"] in ("disclose", "withhold")


def test_group_batch_decide_never_overwrites_individual_decision(client):
    body = _generate_with_sensitive_content(client)
    draft_id = body["drafts"][0]["id"]
    review = client.get(f"/api/disclosure/draft/{draft_id}/review").json()
    assert review["groups"], "expected at least one disclosure group for this fixture"
    group = review["groups"][0]
    item = group["items"][0]

    r1 = client.post(f"/api/disclosure/item/{item['id']}/decide", json={"choice": "withhold", "decided_by": "analyst_a"})
    assert r1.status_code == 200
    assert r1.json()["analyst_choice"] == "withhold"
    assert r1.json()["decision_source"] == "individual"

    r2 = client.post(f"/api/disclosure/draft/{draft_id}/group/{group['group_key']}/decide", json={"choice": "disclose"})
    assert r2.status_code == 200

    review2 = client.get(f"/api/disclosure/draft/{draft_id}/review").json()
    updated_group = next(g for g in review2["groups"] if g["group_key"] == group["group_key"])
    updated_item = next(i for i in updated_group["items"] if i["id"] == item["id"])
    assert updated_item["analyst_choice"] == "withhold"
    assert updated_item["decision_source"] == "individual"


def test_individual_decide_always_overwrites_prior_batch_decision(client):
    body = _generate_with_sensitive_content(client)
    draft_id = body["drafts"][0]["id"]
    review = client.get(f"/api/disclosure/draft/{draft_id}/review").json()
    group = review["groups"][0]
    item = group["items"][0]

    client.post(f"/api/disclosure/draft/{draft_id}/group/{group['group_key']}/decide", json={"choice": "disclose"})
    r = client.post(f"/api/disclosure/item/{item['id']}/decide", json={"choice": "withhold"})
    assert r.status_code == 200
    assert r.json()["analyst_choice"] == "withhold"
    assert r.json()["decision_source"] == "individual"


def test_bulk_accept_all_recommendations_clears_every_pending_item(client):
    body = _generate_with_sensitive_content(client)
    draft_id = body["drafts"][0]["id"]

    r = client.post(f"/api/disclosure/draft/{draft_id}/accept-all-recommendations", json={"decided_by": "analyst"})
    assert r.status_code == 200
    assert r.json()["decision_source"] == "bulk_accept_all"

    review = client.get(f"/api/disclosure/draft/{draft_id}/review").json()
    assert review["pending_count"] == 0
    for group in review["groups"]:
        for item in group["items"]:
            assert item["analyst_choice"] == item["suggested_default"]
            assert item["decision_source"] == "bulk_accept_all"


def test_approval_blocked_409_until_reviewed_then_succeeds_after_bulk_accept(client):
    body = _generate_with_sensitive_content(client)
    draft_id = body["drafts"][0]["id"]

    blocked = client.post("/api/dashboard/approve", json={"draft_id": draft_id})
    assert blocked.status_code == 409

    client.post(f"/api/disclosure/draft/{draft_id}/accept-all-recommendations")

    approved = client.post("/api/dashboard/approve", json={"draft_id": draft_id})
    assert approved.status_code == 200
    assert approved.json()["draft_id"] == draft_id


def test_approval_succeeds_without_409_when_nothing_sensitive_detected(client):
    res = client.post("/api/generate", json={
        "source_text": "Quarterly community newsletter with no sensitive operational details whatsoever.",
        "deliverable_types": ["public_faq"],
    })
    draft_id = res.json()["drafts"][0]["id"]
    approved = client.post("/api/dashboard/approve", json={"draft_id": draft_id})
    assert approved.status_code == 200


def test_withhold_resolution_is_deterministic_no_llm_rewrite(client, fake_llm):
    body = _generate_with_sensitive_content(client)
    draft_id = body["drafts"][0]["id"]
    client.post(f"/api/disclosure/draft/{draft_id}/accept-all-recommendations")

    structured_calls_before = len(fake_llm.structured_calls)
    completions_before = len(fake_llm.complete_calls)

    approved = client.post("/api/dashboard/approve", json={"draft_id": draft_id})
    assert approved.status_code == 200

    assert len(fake_llm.structured_calls) == structured_calls_before
    assert len(fake_llm.complete_calls) == completions_before


def test_aggregation_check_not_wired_into_active_gate(client):
    """Phase 2 correction: aggregation/mosaic-risk checking must not block approval."""
    body = _generate_with_sensitive_content(client)
    draft_id = body["drafts"][0]["id"]
    client.post(f"/api/disclosure/draft/{draft_id}/accept-all-recommendations")
    approved = client.post("/api/dashboard/approve", json={"draft_id": draft_id})
    assert approved.status_code == 200

    review = client.get(f"/api/disclosure/draft/{draft_id}/review").json()
    assert review["aggregation_findings"] == []

    import subprocess
    result = subprocess.run(
        ["grep", "-rl", "--include=*.py", "aggregation_checker", "app/api/"],
        capture_output=True, text=True,
    )
    assert result.stdout.strip() == "", "aggregation_checker must not be wired into any active route"


def test_full_demo_flow_completes_with_zero_unresolved_409s(client):
    """Phase 2.6 acceptance test."""
    body = _generate_with_sensitive_content(client)
    draft_id = body["drafts"][0]["id"]
    client.post(f"/api/disclosure/draft/{draft_id}/accept-all-recommendations")
    approve = client.post("/api/dashboard/approve", json={"draft_id": draft_id})
    assert approve.status_code == 200
    verify = client.get("/api/ledger/verify")
    assert verify.status_code == 200
    assert verify.json()["is_valid"] is True


def test_no_disclosure_endpoint_call_leaves_approval_blocked(client):
    body = _generate_with_sensitive_content(client)
    draft_id = body["drafts"][0]["id"]
    blocked = client.post("/api/dashboard/approve", json={"draft_id": draft_id})
    assert blocked.status_code == 409


# ---------------------------------------------------------------------------
# Phase 3: multimodal ingestion
# ---------------------------------------------------------------------------

def test_docx_ingestion_extracts_paragraph_and_table_text(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_paragraph("An incident occurred at the facility.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Systems affected"
    table.rows[0].cells[1].text = "24"
    path = tmp_path / "test.docx"
    doc.save(str(path))

    result = asyncio.run(ingest_document(file_bytes=path.read_bytes(), filename="test.docx"))
    assert "incident occurred at the facility" in result["raw_text"]
    assert "Systems affected" in result["raw_text"]
    assert result["file_type"] == "docx"


def test_image_ingestion_produces_provenance_note(fake_llm):
    fake_png_bytes = (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
        b"\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc\xf8\xcf\xc0"
        b"\x00\x00\x03\x01\x01\x00\x18\xdd\x8d\xb0\x00\x00\x00\x00IEND\xaeB`\x82"
    )
    result = asyncio.run(ingest_document(file_bytes=fake_png_bytes, filename="entrance.png"))
    assert result["file_type"] == "image"
    assert result["content_provenance_note"] is not None
    assert "machine-generated" in result["content_provenance_note"]
    assert len(result["raw_text"]) > 0


def test_upload_exceeding_max_size_raises_upload_too_large():
    original_limit = settings.MAX_UPLOAD_SIZE_MB
    settings.MAX_UPLOAD_SIZE_MB = 1
    try:
        oversized = b"a" * (2 * 1024 * 1024)
        with pytest.raises(UploadTooLargeError):
            asyncio.run(ingest_document(file_bytes=oversized, filename="big.txt"))
    finally:
        settings.MAX_UPLOAD_SIZE_MB = original_limit


def test_multifile_upload_concatenates_sources():
    result = asyncio.run(ingest_document(
        text="Primary report text about the incident.",
        additional_files=[(b"Supplementary notes on containment.", "notes.txt")],
    ))
    assert "Supplementary notes on containment." in result["raw_text"]
    assert "Primary report text about the incident." in result["raw_text"]


def test_upload_endpoint_returns_413_for_oversized_file(client):
    original_limit = settings.MAX_UPLOAD_SIZE_MB
    settings.MAX_UPLOAD_SIZE_MB = 1
    try:
        big_content = b"x" * (2 * 1024 * 1024)
        res = client.post(
            "/api/generate/upload",
            files={"files": ("big.txt", big_content, "text/plain")},
            data={"deliverable_types": json.dumps(["executive_summary"])},
        )
        assert res.status_code == 413
    finally:
        settings.MAX_UPLOAD_SIZE_MB = original_limit


# ---------------------------------------------------------------------------
# Phase 4: Communication Objective & Content Style auto-inference
# ---------------------------------------------------------------------------

def test_objective_and_style_auto_infer_by_deliverable_type():
    request = GenerateRequest(source_text=SAMPLE_SOURCE, deliverable_types=[DeliverableType.ADVISORY])
    spec = resolver.resolve_spec(DeliverableType.ADVISORY, request)
    assert spec.communication_objective == CommunicationObjective.WARN

    request2 = GenerateRequest(source_text=SAMPLE_SOURCE, deliverable_types=[DeliverableType.PUBLIC_FAQ])
    spec2 = resolver.resolve_spec(DeliverableType.PUBLIC_FAQ, request2)
    assert spec2.content_style == ContentStyle.QA
    assert spec2.communication_objective == CommunicationObjective.INFORM


def test_explicit_objective_and_style_override_auto_inference():
    request = GenerateRequest(
        source_text=SAMPLE_SOURCE,
        deliverable_types=[DeliverableType.ADVISORY],
        communication_objective=CommunicationObjective.REASSURE,
        content_style=ContentStyle.NARRATIVE,
    )
    spec = resolver.resolve_spec(DeliverableType.ADVISORY, request)
    assert spec.communication_objective == CommunicationObjective.REASSURE
    assert spec.content_style == ContentStyle.NARRATIVE


def test_generate_request_accepts_objective_and_style_fields(client):
    res = client.post("/api/generate", json={
        "source_text": SAMPLE_SOURCE,
        "deliverable_types": ["executive_summary"],
        "communication_objective": "Warn",
        "content_style": "Data-led",
    })
    assert res.status_code == 200


# ---------------------------------------------------------------------------
# Phase 5: ledger chain, cross-consistency, verify endpoint
# ---------------------------------------------------------------------------

def test_ledger_chain_verifies_after_three_sequential_approvals(client):
    for i in range(3):
        res = client.post("/api/generate", json={
            "source_text": f"Routine status update number {i} with no sensitive content.",
            "deliverable_types": ["public_faq"],
        })
        draft_id = res.json()["drafts"][0]["id"]
        approve = client.post("/api/dashboard/approve", json={"draft_id": draft_id})
        assert approve.status_code == 200

    verify = client.get("/api/ledger/verify")
    assert verify.status_code == 200
    data = verify.json()
    assert data["is_valid"] is True
    assert data["total_blocks"] == 3


def test_ledger_detects_tampering(client, db_session):
    res = client.post("/api/generate", json={
        "source_text": "Routine status update with no sensitive content.",
        "deliverable_types": ["public_faq"],
    })
    draft_id = res.json()["drafts"][0]["id"]
    client.post("/api/dashboard/approve", json={"draft_id": draft_id})

    from app.db.models import ProvenanceLedgerModel
    block = db_session.query(ProvenanceLedgerModel).order_by(ProvenanceLedgerModel.index.desc()).first()
    block.final_hash = "0" * 64
    db_session.commit()

    verify = client.get("/api/ledger/verify")
    body = verify.json()
    assert body["is_valid"] is False


def test_standalone_verify_endpoint_flags_swapped_number(client):
    res = client.post("/api/verify", json={
        "content_to_verify": "The incident affected 42 systems and was fully contained by evening.",
        "claimed_source": "The incident affected 24 systems at the facility and was fully contained by evening.",
    })
    assert res.status_code == 200
    body = res.json()
    assert body["is_corrupted_detected"] is True
    assert body["trust_score"]["composite_trust_score"] < 70


def test_cross_output_consistency_runs_for_multiselect(client):
    res = client.post("/api/generate", json={
        "source_text": SAMPLE_SOURCE,
        "deliverable_types": ["executive_summary", "linkedin_post"],
    })
    assert res.status_code == 200
    for draft in res.json()["drafts"]:
        assert draft["trust_score"]["consistency_score"] is not None


def test_at_least_one_deliverable_type_required(client):
    res = client.post("/api/generate", json={"source_text": SAMPLE_SOURCE, "deliverable_types": []})
    assert res.status_code == 400


def test_empty_source_text_rejected(client):
    res = client.post("/api/generate", json={"source_text": "   ", "deliverable_types": ["advisory"]})
    assert res.status_code == 400
